"""Render a Markdown document to .docx.

Written because this machine has no pandoc and no LibreOffice, and the summary
carries 16 figures and ~25 tables that have to survive the conversion.

Handles the subset of Markdown the project's documents actually use: ATX
headings, tables with a leading separator row, fenced code, bullet and numbered
lists, block quotes, images with relative paths, horizontal rules, and inline
**bold**, *italic* and `code`. Anything else passes through as body text rather
than being dropped silently.

    python scripts/md_to_docx.py docs/PROJECT_SUMMARY.md -o docs/PROJECT_SUMMARY.docx
"""
from __future__ import annotations

import argparse
import re
from pathlib import Path

import docx
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

INLINE = re.compile(r"(\*\*.+?\*\*|\*[^*]+?\*|`[^`]+?`)")
IMAGE = re.compile(r"^!\[([^\]]*)\]\(([^)]+)\)\s*$")
HEADING = re.compile(r"^(#{1,6})\s+(.*)$")
ULIST = re.compile(r"^(\s*)[-*]\s+(.*)$")
OLIST = re.compile(r"^(\s*)(\d+)\.\s+(.*)$")


def add_runs(par, text: str) -> None:
    """Inline **bold**, *italic* and `code`, with the markers removed."""
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            par.add_run(part[2:-2]).bold = True
        elif part.startswith("`") and part.endswith("`"):
            r = par.add_run(part[1:-1])
            r.font.name = "Consolas"
            r.font.size = Pt(9)
            r.font.color.rgb = RGBColor(0xB0, 0x30, 0x30)
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            par.add_run(part[1:-1]).italic = True
        else:
            par.add_run(part)


def shade(cell, hexcolour: str) -> None:
    el = OxmlElement("w:shd")
    el.set(qn("w:val"), "clear")          # never "solid" -- renders black
    el.set(qn("w:fill"), hexcolour)
    cell._tc.get_or_add_tcPr().append(el)


def split_row(line: str) -> list[str]:
    return [c.strip() for c in line.strip().strip("|").split("|")]


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("src")
    ap.add_argument("-o", "--out", required=True)
    ap.add_argument("--title", default=None)
    a = ap.parse_args(argv)

    src = Path(a.src)
    base = src.parent
    lines = src.read_text().split("\n")
    doc = docx.Document()

    for s in doc.sections:                       # US Letter, 1 inch margins
        s.page_width, s.page_height = Inches(8.5), Inches(11)
        s.left_margin = s.right_margin = Inches(1)
        s.top_margin = s.bottom_margin = Inches(1)
    n = doc.styles["Normal"]
    n.font.name, n.font.size = "Calibri", Pt(10.5)

    if a.title:
        doc.add_heading(a.title, 0)

    i, n_img, n_tab, missing = 0, 0, 0, []
    while i < len(lines):
        ln = lines[i]
        if not ln.strip():
            i += 1
            continue

        m = IMAGE.match(ln.strip())
        if m:
            alt, rel = m.group(1), m.group(2)
            p = (base / rel).resolve()
            if p.exists():
                doc.add_picture(str(p), width=Inches(6.5))
                doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
                n_img += 1
            else:
                missing.append(rel)
                doc.add_paragraph(f"[missing figure: {rel}]")
            if alt:
                cap = doc.add_paragraph()
                cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                r = cap.add_run(alt)
                r.italic, r.font.size = True, Pt(9)
            i += 1
            continue

        m = HEADING.match(ln)
        if m:
            doc.add_heading(m.group(2).strip(), min(len(m.group(1)), 4))
            i += 1
            continue

        if ln.startswith("```"):
            i += 1
            body = []
            while i < len(lines) and not lines[i].startswith("```"):
                body.append(lines[i]); i += 1
            i += 1
            par = doc.add_paragraph()
            par.paragraph_format.left_indent = Inches(0.25)
            par.paragraph_format.space_after = Pt(6)
            r = par.add_run("\n".join(body))
            r.font.name, r.font.size = "Consolas", Pt(8.5)
            continue

        # A table: a pipe row whose successor is the --- separator.
        if ln.strip().startswith("|") and i + 1 < len(lines) and \
                re.match(r"^\s*\|[\s:|-]+\|\s*$", lines[i + 1]):
            head = split_row(ln)
            i += 2
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append(split_row(lines[i])); i += 1
            t = doc.add_table(rows=1, cols=len(head))
            t.style, t.alignment = "Table Grid", WD_TABLE_ALIGNMENT.CENTER
            for c, txt in zip(t.rows[0].cells, head):
                c.text = ""
                add_runs(c.paragraphs[0], txt)
                for r in c.paragraphs[0].runs:
                    r.bold, r.font.size = True, Pt(9)
                shade(c, "DCE6F1")
            for row in rows:
                cells = t.add_row().cells
                for c, txt in zip(cells, row + [""] * (len(head) - len(row))):
                    c.text = ""
                    add_runs(c.paragraphs[0], txt)
                    for r in c.paragraphs[0].runs:
                        r.font.size = Pt(9)
            doc.add_paragraph()
            n_tab += 1
            continue

        if re.match(r"^\s*(---|\*\*\*|___)\s*$", ln):
            doc.add_paragraph()
            i += 1
            continue

        m = ULIST.match(ln)
        if m:
            lvl = min(len(m.group(1)) // 2, 2)
            par = doc.add_paragraph(style="List Bullet" if lvl == 0 else "List Bullet 2")
            add_runs(par, m.group(2))
            i += 1
            continue

        m = OLIST.match(ln)
        if m:
            par = doc.add_paragraph(style="List Number")
            add_runs(par, m.group(3))
            i += 1
            continue

        if ln.startswith(">"):
            par = doc.add_paragraph()
            par.paragraph_format.left_indent = Inches(0.3)
            add_runs(par, ln.lstrip("> ").rstrip())
            for r in par.runs:
                r.italic = True
            i += 1
            continue

        # Body text: join the paragraph's wrapped lines.
        body = [ln.rstrip()]
        i += 1
        while i < len(lines) and lines[i].strip() and not lines[i].startswith(
                ("#", "|", "```", ">", "![")) and not ULIST.match(lines[i]) \
                and not OLIST.match(lines[i]):
            body.append(lines[i].rstrip()); i += 1
        add_runs(doc.add_paragraph(), " ".join(body))

    doc.save(a.out)
    print(f"wrote {a.out}: {len(doc.paragraphs)} paragraphs, {n_tab} tables, "
          f"{n_img} figures embedded")
    if missing:
        print(f"  {len(missing)} figures NOT found:")
        for p in missing:
            print(f"    {p}")


if __name__ == "__main__":
    main()
